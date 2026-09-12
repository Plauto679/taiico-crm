import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from parsers.aarco_prospector_commissions import parse_aarco_prospector_workbook
from parsers.metlife_cobranza import ParsedCobranzaRow
from parsers.sura_cobranza import parse_sura_cobranza_workbook
from services import auth
from services.cobranza_prospectadores import build_preview_workbook, calculate_allocation, consolidate_parsed_rows
from services.prospectadores import normalize_rate, parse_split_prospectors


class ProspectorCommissionTests(unittest.TestCase):
    def test_life_calculation_rounds_each_policy_before_total(self):
        result = calculate_allocation(
            Decimal("116"), Decimal("0.50"), branch="VIDA",
        )
        self.assertEqual(result["commission_base"], Decimal("100.00"))
        self.assertEqual(result["commission_amount"], Decimal("45.50"))
        self.assertEqual(result["vat_amount"], Decimal("7.28"))
        self.assertEqual(result["total_amount"], Decimal("52.78"))

    def test_negative_reversal_reduces_commission_and_vat(self):
        result = calculate_allocation(
            Decimal("-116"), Decimal("0.50"), branch="VIDA",
        )
        self.assertEqual(result["commission_amount"], Decimal("-45.50"))
        self.assertEqual(result["vat_amount"], Decimal("-7.28"))
        self.assertEqual(result["total_amount"], Decimal("-52.78"))

    def test_legacy_whole_and_fractional_rates_are_normalized(self):
        self.assertEqual(normalize_rate("50"), Decimal("0.500000"))
        self.assertEqual(normalize_rate("0.5"), Decimal("0.500000"))

    def test_explicit_split_prospectors_are_preserved(self):
        rows = parse_split_prospectors("50% Norma, 10% Pamela, 40% Paola")
        self.assertEqual([name for name, _ in rows], ["Norma", "Pamela", "Paola"])
        self.assertEqual(sum((rate for _, rate in rows), Decimal("0")), Decimal("1.000000"))

    def test_new_modules_are_available_in_access_configuration(self):
        config = auth.access_modules_configuration()
        modules = {item["key"]: item for item in config["modules"]}
        self.assertEqual(modules["prospectadores"]["label"], "Prospectadores")
        self.assertEqual(modules["cobranza_prospectadores"]["label"], "Cobranza para prospectadores")

    def test_metlife_rows_are_consolidated_by_policy_and_payment_date(self):
        rows = []
        for index, amount in enumerate(("100.25", "25.75"), start=2):
            rows.append(ParsedCobranzaRow(
                parser_name="test", parser_version="1", sheet_name="Vida",
                row_number=index, row_hash=str(index), source_payload={}, issues=[],
                normalized_payload={
                    "insurer_id": "metlife", "product_branch": "VIDA",
                    "policy_number": "008123", "payment_date": date(2026, 8, 3),
                    "net_commission_amount": Decimal(amount),
                },
            ))
        consolidated = consolidate_parsed_rows("metlife", rows)
        self.assertEqual(len(consolidated), 1)
        self.assertEqual(consolidated[0]["source_commission"], "126.0000")
        self.assertEqual(consolidated[0]["row_count"], 2)

    def test_sura_consolidated_format_accepts_missing_currency_column(self):
        headers = [
            "FECHA DE CORTE", "OFICINA", "RAMO", "POLIZA", "RECIBO",
            "SERIE_RECIBO", "AGENTE", "INIVIGENCIA", "FECHAAPLICACION",
            "ASEGURADO", "PRIMATOTAL", "PRIMANETA", "COMNETA",
            "%COMISION", "COMDERECHO", "TOTAL",
        ]
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as workbook:
            pd.DataFrame([[
                "31/08/2026", 1, 606, 19663, 16354889, "01/01", "", "25/07/2026",
                "31/08/2026", "Cliente", 2617.95, 2256.86, 338.54, 15, 0, 392.71,
            ]], columns=headers).to_excel(workbook.name, index=False)
            rows, issues = parse_sura_cobranza_workbook(workbook.name)
        self.assertEqual(issues, [])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].normalized_payload["currency"], "MXN")

    def test_aarco_parser_uses_applied_commission_and_ignores_legacy_percentage(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as workbook:
            pd.DataFrame([{
                "CIA": "AXA", "NUM_POL": "0862932H", "CLIENTE": "Cliente",
                "F_COBRO": "2026-08-27", "PRIMA_NETA_MN": 862.39,
                "COM_APL_MN": 32.38, "%\nCOMISION \nPROSPECTADOR": 0,
                "$\nCOMISION \nPROSPECTADOR": 0,
            }]).to_excel(workbook.name, index=False)
            rows, issues = parse_aarco_prospector_workbook(workbook.name)
        self.assertEqual(issues, [])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].normalized_payload["source_commission_amount"], Decimal("32.38"))

    def test_preview_export_contains_all_rows_and_keeps_amounts_numeric(self):
        rows = [
            {
                "insurer_id": "metlife", "policy_number": "00123", "receipt_number": "R-1",
                "branch": "VIDA", "movement_date": "2026-08-01", "currency": "MXN",
                "source_commission": "116.00", "row_count": 2, "matching_hints": {},
                "status": "listo", "exception_reason": None,
                "allocations": [{
                    "prospector_name": "ANA PRUEBA", "commission_rate": "0.5",
                    "calculation": {"total_amount": "52.78"},
                }],
            },
            {
                "insurer_id": "metlife", "policy_number": "00999", "receipt_number": "",
                "branch": "GMM", "movement_date": "2026-08-02", "currency": "MXN",
                "source_commission": "80.25", "row_count": 1,
                "matching_hints": {"office_code": "01", "branch_code": "606"},
                "status": "pendiente_asignacion", "allocations": [],
                "exception_reason": "La póliza no tiene un prospectador asignado",
            },
        ]
        workbook = load_workbook(BytesIO(build_preview_workbook("metlife", "agosto.xlsx", rows)))
        sheet = workbook["Vista previa"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["A1"].value, "Aseguradora")
        self.assertEqual(sheet["B3"].value, "00999")
        self.assertIsInstance(sheet["G2"].value, (int, float))
        self.assertEqual(sheet["G2"].value, 116)
        self.assertEqual(sheet["L3"].value, "La póliza no tiene un prospectador asignado")
        self.assertEqual(workbook["Resumen"]["B6"].value, 1)


if __name__ == "__main__":
    unittest.main()
