from datetime import date, datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from io import BytesIO
import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from database import Base
from services import agent_statements as s

ACTOR=SimpleNamespace(username='test@example.com')

@pytest.fixture
def store(tmp_path,monkeypatch):
    engine=create_engine('sqlite:///'+str(tmp_path/'test.db'))
    Base.metadata.create_all(engine)
    factory=sessionmaker(bind=engine,autoflush=False)
    monkeypatch.setattr(s,'SessionLocal',factory)
    with factory() as db:
        db.add(s.Prospector(id='p',name='Prospectador de prueba',normalized_name='PROSPECTADOR DE PRUEBA',email='p@example.com',created_by='test'))
        db.add(s.Prospector(id='q',name='Sin comisiones',normalized_name='SIN COMISIONES',created_by='test'))
        for i,m in enumerate((8,9,10)):
            db.add(s.ProspectorCommissionPeriod(id=str(m),month=date(2026,m,1),created_by='test'))
            for source in ('metlife','sura','aarco'):
                db.add(s.ProspectorCommissionBatch(id=f'{m}-{source}',period_id=str(m),source=source,filename='test.xlsx',file_hash=f'{m}-{source}',created_by='test'))
            db.add(s.ProspectorCommissionLine(id=f'l{m}',batch_id=f'{m}-metlife',source_key=f'l{m}',policy_number='123',insurer_id='metlife',source_commission=100,status='listo'))
            db.add(s.ProspectorCommissionAllocation(id=f'a{m}',line_id=f'l{m}',prospector_id='p',commission_rate='.5',utility_coefficient='.91',vat_rate='.16',life_divisor=1,commission_amount=100,vat_amount=0,total_amount=100,calculation_json={}))
        db.commit()
    yield factory
    engine.dispose()


def state(month):
    return next(x for x in s.workspace()['statements'] if x['month'].startswith(f'2026-{month:02}') and x['prospector_id']=='p')


def generate(month):
    return s.generate(s.GenerateInput(period_id=str(month)),ACTOR)


def test_generate_idempotent_adjustments_and_three_month_carry(store):
    assert generate(8)['created']==1
    a=state(8)
    s.edit(a['id'],s.EditInput(revision=a['revision'],recipient='p@example.com',adjustments=[s.Adjustment(concept='Saldo inicial',amount='50')]),ACTOR)
    assert generate(8)['created']==0
    assert state(8)['total']=='158.00'
    generate(9);generate(10)
    assert state(9)['total']=='258.00'
    assert state(10)['total']=='358.00' # not 100+150+250
    c=state(10)
    s.set_status(c['id'],s.StatusInput(revision=c['revision'],status='pagado'),ACTOR)
    assert all(x['status']=='pagado' for x in s.workspace()['statements'])
    assert generate(10)['updated']==0


def test_correction_updates_unpaid_later_statements_and_preserves_paid(store):
    for m in (8,9,10):generate(m)
    a=state(8);b=state(9)
    with store() as db:
        db.get(s.AgentStatement,b['id']).status='enviado';db.commit()
    s.edit(a['id'],s.EditInput(revision=a['revision'],adjustments=[s.Adjustment(concept='Corrección',amount='-25')]),ACTOR)
    assert state(9)['status']=='pendiente'
    assert state(10)['total']=='271.00'
    a=state(8)
    s.set_status(a['id'],s.StatusInput(revision=a['revision'],status='pagado'),ACTOR)
    assert state(10)['total']=='200.00'
    with pytest.raises(HTTPException):
        s.edit(a['id'],s.EditInput(revision=state(8)['revision'],adjustments=[]),ACTOR)


def test_missing_sources_exceptions_and_stale_edit_are_blocked(store):
    with store() as db:
        db.get(s.ProspectorCommissionLine,'l8').status='pendiente_asignacion';db.commit()
    with pytest.raises(HTTPException): generate(8)
    with store() as db:
        db.get(s.ProspectorCommissionLine,'l8').status='listo';db.commit()
    generate(8);a=state(8)
    with pytest.raises(HTTPException):s.edit(a['id'],s.EditInput(revision=0),ACTOR)
    with store() as db:
        db.delete(db.get(s.ProspectorCommissionBatch,'9-sura'));db.commit()
    with pytest.raises(HTTPException):generate(9)


def test_create_manual_without_commissions_and_zero_amount(store):
    s.generate(s.GenerateInput(period_id='8',prospector_id='q'),ACTOR)
    row=next(x for x in s.workspace()['statements'] if x['prospector_id']=='q')
    assert row['total']=='0.00'
    s.edit(row['id'],s.EditInput(revision=row['revision'],adjustments=[s.Adjustment(concept='Pendiente anterior',amount='250')]),ACTOR)
    generate(9)
    row=next(x for x in s.workspace()['statements'] if x['prospector_id']=='q' and x['period_id']=='9')
    assert row['total']=='290.00'


def test_sending_is_individual_and_does_not_resend_without_confirmation(store):
    generate(8);a=state(8)
    settings={'sender':'sender@example.com','host':'test','port':587,'user':'test','password':'test','use_starttls':True}
    with patch.object(s,'smtp_settings_for',return_value=settings),patch.object(s.smtplib,'SMTP') as smtp:
        assert s.send(a['id'],s.SendInput(revision=a['revision']),ACTOR)=={'sent':True}
        message=smtp.return_value.__enter__.return_value.send_message.call_args.args[0]
        assert str(message['To'])=='p@example.com'
        assert str(message['Cc'])=='veronica.alfaro@taiico.com'
        body=message.get_body(preferencelist=('plain',)).get_content()
        assert 'Importe Sin IVA: $100.00' in body
        assert 'IVA: $0.00' in body
        assert 'Total: $100.00.' in body
        assert 'Recuerda enviar la factura correspondiente a veronica.alfaro@taiico.com' in body
        assert len(list(message.iter_attachments()))==2
        assert state(8)['status']=='enviado'
        with pytest.raises(HTTPException):s.send(a['id'],s.SendInput(revision=state(8)['revision']),ACTOR)
        assert smtp.call_count==1


def test_failed_delivery_requires_manual_review_and_inflight_blocks_edits(store):
    generate(8);a=state(8)
    settings={'sender':'sender@example.com','host':'test','port':587,'user':'test','password':'test','use_starttls':True}
    with patch.object(s,'smtp_settings_for',return_value=settings),patch.object(s.smtplib,'SMTP',side_effect=OSError('simulated failure')) as smtp:
        with pytest.raises(HTTPException):s.send(a['id'],s.SendInput(revision=a['revision']),ACTOR)
        row=state(8);assert row['delivery_state']=='uncertain';assert row['status']=='pendiente'
        with pytest.raises(HTTPException):s.send(a['id'],s.SendInput(revision=row['revision'],resend=True),ACTOR)
        assert smtp.call_count==1
    s.resolve_delivery(a['id'],s.DeliveryResolution(revision=row['revision'],delivered=False),ACTOR)
    with store() as db:
        db.get(s.AgentStatement,a['id']).delivery_state='sending';db.commit()
    with pytest.raises(HTTPException):s.edit(a['id'],s.EditInput(revision=state(8)['revision']),ACTOR)


def test_pdf_excel_content_and_formula_escaping(store):
    from openpyxl import load_workbook
    generate(8);a=state(8)
    s.edit(a['id'],s.EditInput(revision=a['revision'],adjustments=[s.Adjustment(concept='=SUM(A1)',amount='10')]),ACTOR)
    data=s.detail(a['id'])
    data['lines'][0]['rate']='0.509999'
    original_amounts={k:data['lines'][0][k] for k in ('commission','vat','total')}
    assert s.pdf_bytes(data).startswith(b'%PDF')
    workbook=load_workbook(BytesIO(s.excel_bytes(data)))
    assert not any(cell.data_type=='f' for row in workbook.active for cell in row)
    assert workbook.active['D7'].value==111.6
    assert workbook.active['E10'].value==.5
    assert workbook.active['E10'].number_format=='0%'
    assert workbook.active['C10'].value=='MetLife'
    assert len(workbook.active._images)==1
    assert {k:data['lines'][0][k] for k in original_amounts}==original_amounts
    assert s.display_percentage('0.509999')==50
    assert s.display_percentage('-0.509999')==-50


def test_email_added_after_generation_is_used_in_list_detail_and_send(store):
    with store() as db:
        db.get(s.Prospector,'p').email=None;db.commit()
    generate(8);a=state(8)
    assert a['recipient']==''
    with store() as db:
        db.get(s.Prospector,'p').email='new@example.com';db.commit()
    assert state(8)['recipient']=='new@example.com'
    assert s.detail(a['id'])['recipient']=='new@example.com'
    settings={'sender':'sender@example.com','host':'test','port':587,'user':'test','password':'test','use_starttls':True}
    with patch.object(s,'smtp_settings_for',return_value=settings),patch.object(s.smtplib,'SMTP') as smtp:
        s.send(a['id'],s.SendInput(revision=a['revision']),ACTOR)
        message=smtp.return_value.__enter__.return_value.send_message.call_args.args[0]
        assert str(message['To'])=='new@example.com'
    with store() as db:
        assert db.get(s.AgentStatement,a['id']).recipient=='new@example.com'


def test_statement_specific_email_is_preserved(store):
    generate(8);a=state(8)
    s.edit(a['id'],s.EditInput(revision=a['revision'],recipient='custom@example.com'),ACTOR)
    with store() as db:
        db.get(s.Prospector,'p').email='changed@example.com';db.commit()
    assert state(8)['recipient']=='custom@example.com'
    assert s.detail(a['id'])['recipient']=='custom@example.com'


def test_adjustment_vat_and_carry_are_not_taxed_twice(store):
    generate(8);a=state(8)
    s.edit(a['id'],s.EditInput(revision=a['revision'],adjustments=[s.Adjustment(concept='Saldo previo',amount='100')]),ACTOR)
    first=s.detail(a['id'])
    assert first['adjustments'][0]['amount']=='100'
    assert first['adjustments'][0]['vat']=='16.00'
    assert first['total']=='216.00'
    generate(9);generate(10)
    third=state(10)
    data=s.detail(third['id'])
    matrix=s.summary_matrix(data)
    assert matrix[2][1:]==(s.Decimal('300.00'),s.Decimal('16.00'),s.Decimal('316.00'))
    assert matrix[-1][-1]==s.Decimal(third['total'])==s.Decimal('416.00')


def test_legacy_final_adjustments_keep_payable_amount(store):
    generate(8);a=state(8)
    with store() as db:
        row=db.get(s.AgentStatement,a['id'])
        row.adjustments=[{'concept':'Anterior','amount':'116.00'}]
        row.own_amount=216;row.total_amount=216
        db.commit()
    data=s.detail(a['id'])
    assert data['adjustments'][0]['amount']=='100.00'
    assert data['adjustments'][0]['vat']=='16.00'
    assert s.summary_matrix(data)[-1][-1]==s.Decimal('216.00')


def test_zero_commissions_do_not_create_automatic_statement(store):
    with store() as db:
        allocation=db.get(s.ProspectorCommissionAllocation,'a8')
        allocation.commission_rate=0;allocation.commission_amount=0
        allocation.vat_amount=0;allocation.total_amount=0
        db.commit()
    assert generate(8)['created']==0
    assert not s.workspace()['statements']
    s.generate(s.GenerateInput(period_id='8',prospector_id='p'),ACTOR)
    assert len(s.workspace()['statements'])==1


def test_zero_rate_lines_are_absent_from_detail_and_exports(store):
    from openpyxl import load_workbook
    generate(8);a=state(8)
    with store() as db:
        row=db.get(s.AgentStatement,a['id'])
        row.lines=[*row.lines, {**row.lines[0], 'policy':'ZERO-POLICY', 'rate':'0', 'commission':'0','vat':'0','total':'0'}]
        db.commit()
    data=s.detail(a['id'])
    assert all(x['policy']!='ZERO-POLICY' for x in data['lines'])
    workbook=load_workbook(BytesIO(s.excel_bytes(data)))
    assert not any(cell.value=='ZERO-POLICY' for row in workbook.active for cell in row)
    assert data['total']=='100.00'


@pytest.mark.parametrize('status,resend',[('pendiente',False),('enviado',True)])
def test_inactive_prospector_cannot_send_even_with_old_screen(store,status,resend):
    generate(8);a=state(8)
    with store() as db:
        db.get(s.Prospector,'p').is_active=False
        db.get(s.AgentStatement,a['id']).status=status
        db.commit()
    assert s.detail(a['id'])['prospector_active'] is False
    settings={'sender':'sender@example.com','host':'test','port':587,'user':'test','password':'test','use_starttls':True}
    with patch.object(s,'smtp_settings_for',return_value=settings),patch.object(s.smtplib,'SMTP') as smtp:
        with pytest.raises(HTTPException) as err:
            s.send(a['id'],s.SendInput(revision=a['revision'],resend=resend),ACTOR)
        assert err.value.status_code==409
        assert 'inactivo' in err.value.detail
        smtp.assert_not_called()
    assert state(8)['status']==status
    assert state(8)['delivery_state']==''
    with store() as db:
        db.get(s.Prospector,'p').is_active=True;db.commit()
    assert s.detail(a['id'])['prospector_active'] is True


def test_manual_policy_conversion_regeneration_exports_and_carry(store):
    from openpyxl import load_workbook
    generate(8);a=state(8)
    s.edit(a['id'],s.EditInput(revision=a['revision'],adjustments=[s.Adjustment(concept='16784T03 AXA',amount='2095.53')]),ACTOR)
    previous=state(8)
    item=s.ManualCommission(policy='16784T03',date='2026-08-15',insurer='AXA',branch='DANOS',commission='2095.53')
    data=s.edit(a['id'],s.EditInput(revision=previous['revision'],manual_commissions=[item],adjustments=[]),ACTOR)
    assert data['total']==previous['total']
    assert data['adjustment_total']=='0.00'
    assert data['commission']=='2530.81'
    manual=next(x for x in data['lines'] if x.get('manual'))
    assert manual['vat']=='335.28'
    assert manual['rate'] is None
    generate(8)
    assert next(x for x in s.detail(a['id'])['lines'] if x.get('manual'))==manual
    assert s.pdf_bytes(s.detail(a['id'])).startswith(b'%PDF')
    ws=load_workbook(BytesIO(s.excel_bytes(s.detail(a['id'])))).active
    assert any(c.value=='16784T03' for row in ws for c in row)
    generate(9)
    assert state(9)['total']=='2630.81'
    # Omitted manual list preserves manual records for older clients.
    s.edit(a['id'],s.EditInput(revision=state(8)['revision']),ACTOR)
    assert len([x for x in s.detail(a['id'])['lines'] if x.get('manual')])==1
    # Explicit deletion removes the amount and updates the carry exactly once.
    s.edit(a['id'],s.EditInput(revision=state(8)['revision'],manual_commissions=[]),ACTOR)
    assert state(9)['total']=='200.00'


def test_manual_policy_validation_and_send(store):
    generate(8);a=state(8)
    with pytest.raises(HTTPException):
        s.edit(a['id'],s.EditInput(revision=a['revision'],manual_commissions=[
            s.ManualCommission(policy='TEST',date='2026-09-01',insurer='AXA',branch='DANOS',commission='10')]),ACTOR)
    assert state(8)['revision']==a['revision']
    data=s.edit(a['id'],s.EditInput(revision=a['revision'],manual_commissions=[
        s.ManualCommission(policy='TEST',date='2026-08-01',insurer='AXA',branch='DANOS',percentage='80',commission='10')]),ACTOR)
    assert data['total']=='111.60' # percentage is informative, already applied by the operator
    with patch.object(s,'smtp_settings_for',return_value={'sender':'sender@example.com','host':'test','port':587,'user':'test','password':'test','use_starttls':True}),patch.object(s.smtplib,'SMTP') as smtp:
        s.send(a['id'],s.SendInput(revision=data['revision']),ACTOR)
        smtp.return_value.__enter__.return_value.send_message.assert_called_once()
